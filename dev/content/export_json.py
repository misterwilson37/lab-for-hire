"""Convert the content workbook into the game's JSON shape.
This is the bootstrap path only: once the Apps Script feed exists, it produces this same JSON
from the live Google Sheet, and this script is retired."""
import json, sys
from openpyxl import load_workbook

src = sys.argv[1]
out = sys.argv[2]
wb = load_workbook(src, data_only=True)

def rows(tab):
    ws = wb[tab]
    it = ws.iter_rows(values_only=True)
    next(it)
    return [r for r in it if r and r[0] not in (None, "")]

def s(v):
    return "" if v is None else str(v).strip()

def n(v):
    if v in (None, ""):
        return None
    return float(v) if not isinstance(v, (int, float)) else v

def axis(label, unit, mn, mx, step, scale):
    return {"label": s(label), "unit": s(unit), "min": n(mn), "max": n(mx), "step": n(step),
            "scale": (s(scale) or "Linear")}

C = {"schema": "0.3.0", "source": src.split("/")[-1]}
C["cases"] = [{"id": s(r[0]), "status": s(r[1]), "order": int(r[2] or 99), "title": s(r[3]), "client": s(r[4]),
               "pitch": s(r[5]), "win": s(r[6]), "requires": s(r[7]), "pay": int(r[8] or 0), "standard": s(r[9])} for r in rows("Cases")]
C["items"] = [{"id": s(r[0]), "name": s(r[1]), "slot": s(r[2]), "price": int(r[3] or 0), "kind": s(r[4]), "unlocks": s(r[5]),
               "rank": s(r[6]), "sprite": s(r[7]), "desc": s(r[8])} for r in rows("Items")]
C["visuals"] = [{"id": s(r[0]), "caseId": s(r[1]), "label": s(r[2]), "type": s(r[3]), "title": s(r[4]),
                 "build": s(r[5]) == "Yes", "x": axis(*r[6:12]), "y": axis(*r[12:18])} for r in rows("Visuals")]
C["data"] = [{"visualId": s(r[0]), "series": s(r[1]), "x": n(r[2]), "y": n(r[3]), "label": s(r[4])}
             for r in rows("Data")]
C["tables"] = [{"visualId": s(r[0]), "row": int(r[1]), "cells": [s(c) for c in r[2:8]]} for r in rows("Tables")]
C["questions"] = [{"id": s(r[0]), "caseId": s(r[1]), "order": int(r[2] or 0), "role": s(r[3]),
                   "forTag": s(r[4]), "skill": s(r[5]), "type": s(r[6]),
                   "visuals": [v.strip() for v in s(r[7]).split(",") if v.strip()],
                   "prompt": s(r[8]), "why": s(r[9])} for r in rows("Questions")]
C["options"] = [{"qid": s(r[0]), "letter": s(r[1]), "text": s(r[2]), "correct": s(r[3]) == "Yes", "tag": s(r[4])}
                for r in rows("Options")]
C["tags"] = [{"id": s(r[0]), "kind": s(r[1]), "title": s(r[2]), "explain": s(r[3]), "teacherNote": s(r[4]),
              "drill": s(r[5])} for r in rows("SideMissions")]
json.dump(C, open(out, "w"), ensure_ascii=False, separators=(",", ":"))
print("exported", {k: len(v) for k, v in C.items() if isinstance(v, list)})
