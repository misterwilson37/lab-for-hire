"""One-time SEED for the Items tab. After Jake pastes it into the content Sheet, the Sheet is the
source of truth (Rule 9) and this file is history. The bake reads an .xlsx export of the Sheet."""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

HEAD = ["ItemID", "Name", "Slot", "Price", "Kind", "Unlocks", "Rank required", "Sprite", "Description"]
R = [
 # ---- starter lab: price 0 = owned from the start, placed in row order ----
 ("desk", "Starter Desk", "bench", 0, "Cosmetic", "", 0, "desk", "A plain wooden desk. Every lab starts somewhere."),
 ("chair-desk", "Desk Chair", "chair", 0, "Cosmetic", "", 0, "chairdesk", "Rolls, spins, and squeaks a little."),
 ("shelf-open", "Open Shelf", "shelf", 0, "Cosmetic", "", 0, "bookcaseopen", "Room for binders, samples, and snacks."),
 ("plant-potted", "Potted Plant", "plant", 0, "Cosmetic", "", 0, "pottedplant", "Needs light and water, like the plants in your cases."),
 ("beaker", "Beaker", "bench-item", 0, "Equipment", "", 0, "beaker", "Holds and pours liquids. The marks on the side are for estimating."),
 ("flask", "Flask", "bench-item", 0, "Equipment", "", 0, "flask", "The narrow neck means fewer splashes when you swirl."),
 ("poster-line", "Line Graph Poster", "poster", 0, "Cosmetic", "", 0, "poster-line", "A line graph shows how something changes."),
 ("wall-white", "Lab White", "paint-wall", 0, "Cosmetic", "", 0, "#F7F5FA", "Clean and bright."),
 ("floor-lilac", "Lilac Gray", "paint-floor", 0, "Cosmetic", "", 0, "#D8D3E3", "The floor every lab starts with."),
 ("trim-purple", "Purple", "paint-trim", 0, "Cosmetic", "", 0, "#6547CF", "The classic Lab for Hire purple."),
 ("tiles-plain", "Plain", "paint-tiles", 0, "Cosmetic", "", 0, "none", "No pattern."),
 # ---- equipment: owning it unlocks case types ----
 ("microscope", "Microscope", "bench-item", 150, "Equipment", "cells", 0, "microscope", "Magnifies things too small to see. Unlocks cell cases (7.LS1.1, 7.LS1.2)."),
 ("grow-cabinet", "Grow Cabinet", "machine", 200, "Equipment", "plant-growth", 1, "grow-cabinet", "A mini greenhouse with its own grow light. Unlocks plant growth cases (7.LS1.4)."),
 ("materials-tester", "Materials Tester", "machine", 250, "Equipment", "materials", 1, "materials-tester", "Stretches and squeezes samples so you can compare them fairly. Unlocks medical design cases (7.ETS1.1)."),
 ("calorimeter", "Calorimeter", "bench-item", 200, "Equipment", "reactions", 2, "calorimeter", "Measures heat given off or taken in. Unlocks reaction cases (7.PS3.1)."),
 ("tube-rack", "Test Tube Rack", "bench-item", 40, "Equipment", "", 0, "tube-rack", "Three samples side by side, easy to compare."),
 # ---- cosmetics ----
 ("lab-bench", "Lab Bench", "bench", 180, "Cosmetic", "", 1, "lab-bench", "A real lab bench with a tough dark top."),
 ("generator", "Power Generator", "machine", 120, "Cosmetic", "", 0, "generator", "Hums quietly. Mostly for looks."),
 ("rug-round", "Round Rug", "rug", 40, "Cosmetic", "", 0, "ruground", "Soft, round, and a little bit fancy."),
 ("rug-long", "Long Rug", "rug", 45, "Cosmetic", "", 0, "rugrectangle", "Covers more floor."),
 ("poster-bars", "Bar Graph Poster", "poster", 30, "Cosmetic", "", 0, "poster-bars", "A bar graph compares groups."),
 ("poster-scatter", "Scatter Plot Poster", "poster", 30, "Cosmetic", "", 1, "poster-scatter", "A scatter plot helps you spot a trend."),
 ("shelf-closed", "Bookcase", "shelf", 60, "Cosmetic", "", 0, "bookcaseclosed", "Solid sides keep your stuff from falling out."),
 ("chair-cushion", "Cushion Chair", "chair", 50, "Cosmetic", "", 0, "chaircushion", "Less rolling, more comfy."),
 ("books", "Stack of Books", "shelf-top", 20, "Cosmetic", "", 0, "books", "Field guides and lab manuals."),
 ("pet-cat", "Lab Cat", "pet", 80, "Cosmetic", "", 0, "pet-cat", "Supervises every experiment."),
 ("pet-dog", "Lab Dog", "pet", 80, "Cosmetic", "", 0, "pet-dog", "Very excited about your data."),
 ("pet-bunny", "Lab Bunny", "pet", 80, "Cosmetic", "", 1, "pet-bunny", "Twitchy nose, big ears."),
 ("wall-lavender", "Lavender", "paint-wall", 40, "Cosmetic", "", 0, "#F3EEFF", "A soft purple glow."),
 ("wall-mint", "Mint", "paint-wall", 40, "Cosmetic", "", 0, "#E1F3EA", "Cool and calm."),
 ("floor-periwinkle", "Periwinkle", "paint-floor", 40, "Cosmetic", "", 0, "#C9D6F6", "A floor with a little blue in it."),
 ("trim-teal", "Teal", "paint-trim", 25, "Cosmetic", "", 0, "#1F8FBF", "Trim with a splash of teal."),
 ("tiles-checker", "Checker", "paint-tiles", 50, "Cosmetic", "", 0, "checker", "Big checkerboard tiles."),
 ("tiles-stripes", "Stripes", "paint-tiles", 50, "Cosmetic", "", 0, "stripes", "Wide floor stripes."),
 # ---- theme pack: Frankenstein ----
 ("fk-candles", "Candle Cluster", "shelf-top", 30, "Theme", "", 0, "candles", "Frankenstein pack. Spooky light for late-night experiments."),
 ("fk-lantern", "Candle Lantern", "lantern", 40, "Theme", "", 0, "lantern", "Frankenstein pack. Lights up the front corner."),
 ("fk-pumpkin", "Jack-o'-lantern", "lantern", 35, "Theme", "", 0, "pumpkin", "Frankenstein pack. Carved and grinning."),
 ("fk-coil", "Spark Coil", "machine", 150, "Theme", "", 1, "spark-coil", "Frankenstein pack. Big sparks, no real science. Just for looks."),
 ("fk-walls", "Stormy", "paint-wall", 45, "Theme", "", 0, "#9FAAA2", "Frankenstein pack. A gloomy, stormy-night gray."),
]
SLOTS = ["bench", "bench-item", "chair", "shelf", "shelf-top", "plant", "machine", "rug", "poster", "pet", "lantern", "paint-wall", "paint-floor", "paint-trim", "paint-tiles"]

def build(path):
    wb = Workbook(); ws = wb.active; ws.title = "Items"
    ws.append(HEAD)
    for r in R: ws.append(list(r))
    bold = Font(name="Arial", bold=True, color="FFFFFF"); norm = Font(name="Arial")
    for c in ws[1]: c.font = bold; c.fill = PatternFill("solid", fgColor="6547CF"); c.alignment = Alignment(vertical="center")
    for row in ws.iter_rows(min_row=2):
        for c in row: c.font = norm; c.alignment = Alignment(vertical="top", wrap_text=(c.column == 9))
    for col, w in zip("ABCDEFGHI", [18, 22, 13, 8, 12, 14, 14, 17, 70]): ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    n = len(R) + 1
    dv1 = DataValidation(type="list", formula1='"' + ",".join(SLOTS) + '"', allow_blank=False); dv1.add(f"C2:C{n+60}")
    dv2 = DataValidation(type="list", formula1='"Equipment,Cosmetic,Theme"', allow_blank=False); dv2.add(f"E2:E{n+60}")
    dv3 = DataValidation(type="whole", operator="between", formula1="0", formula2="100000"); dv3.add(f"D2:D{n+60}")
    dv4 = DataValidation(type="whole", operator="between", formula1="0", formula2="20"); dv4.add(f"G2:G{n+60}")
    for d in (dv1, dv2, dv3, dv4): ws.add_data_validation(d)

    lg = wb.create_sheet("Items legend")
    rows = [
     ("Column", "What goes here"),
     ("ItemID", "Unique, lowercase, never renamed once students own it (saved labs store this ID)."),
     ("Name", "What students see."),
     ("Slot", "Where it goes: " + ", ".join(SLOTS) + ". bench-item fills any of 3 bench spots; poster fills either wall. paint-* rows are recolors."),
     ("Price", "Lab funds. 0 = part of the starter lab (owned from the start, placed in row order)."),
     ("Kind", "Equipment (can unlock case types), Cosmetic, or Theme (a seasonal pack)."),
     ("Unlocks", "Optional case type this item unlocks when OWNED (placing it isn't required): cells, plant-growth, materials, reactions."),
     ("Rank required", "Rank number the student must reach before buying. 0 = anyone."),
     ("Sprite", "Art recipe ID from bake/scene.js (RECIPES). For paint rows: a #hex color, or none / checker / stripes for floor patterns."),
     ("Description", "One or two short sentences. Equipment names the case type and standard it unlocks."),
     ("", ""),
     ("Build note", "Sprites and lab-manifest.json are generated from this tab by bake/bake.py. Never hand-edit them. A new Sprite recipe needs code in scene.js; a new item using an existing recipe needs only a row here and a re-bake."),
    ]
    for r in rows: lg.append(list(r))
    for c in lg[1]: c.font = bold; c.fill = PatternFill("solid", fgColor="6547CF")
    for row in lg.iter_rows(min_row=2):
        for c in row: c.font = Font(name="Arial", bold=(c.column == 1)); c.alignment = Alignment(wrap_text=True, vertical="top")
    lg.column_dimensions["A"].width = 16; lg.column_dimensions["B"].width = 100
    wb.save(path)

if __name__ == "__main__":
    build(sys.argv[1] if len(sys.argv) > 1 else "LabForHire_Items_v0.1.0.xlsx")
