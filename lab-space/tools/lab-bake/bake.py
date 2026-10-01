"""Lab Space sprite bake.  Items tab (.xlsx export of the content Sheet) -> lab-sprites/*.png + lab-manifest.json + lab-items.json

Every item is rendered ALONE, IN ITS SLOT, through the ONE locked camera and lighting rig (scene.js), to a
full-frame transparent PNG, then cropped to its bounding box with the offset kept. Because every sprite shares
the room's frame, the game just stacks images in slot order: no anchor math.

usage: python3 bake.py <Items.xlsx> <out_dir>
Outputs are build artifacts (Rule 9): never hand-edit them; change the Sheet or scene.js and re-bake.
"""
import sys, os, io, json, base64, re, hashlib
from openpyxl import load_workbook
from PIL import Image
from harness import serve, GL_ARGS
from playwright.sync_api import sync_playwright

MANIFEST_VERSION = "0.3.0"
KINDS = ("Equipment", "Cosmetic", "Theme", "Award")
THUMB = 112  # store thumbnails: 56 CSS px at 2x
HEAD = ["ItemID", "Name", "Slot", "Price", "Kind", "Unlocks", "Rank required", "Sprite", "Description"]
KEYS = ["id", "name", "slot", "price", "kind", "unlocks", "rank", "sprite", "description"]
PAINT = {"paint-wall": "wall", "paint-floor": "floor", "paint-trim": "trim", "paint-tiles": "tiles"}
PATTERNS = {"none", "checker", "stripes"}

def read_items(path):
    wb = load_workbook(path, data_only=True)
    if "Items" not in wb.sheetnames: sys.exit("no Items tab in " + path)
    ws = wb["Items"]; rows = list(ws.iter_rows(values_only=True))
    head = [str(h).strip() if h is not None else "" for h in rows[0]]
    if head[:len(HEAD)] != HEAD: sys.exit(f"Items header must be {HEAD}, got {head}")
    items = []
    for r in rows[1:]:
        if not r or r[0] in (None, ""): continue
        it = dict(zip(KEYS, r))
        for k in ("id", "name", "slot", "kind", "unlocks", "sprite", "description"): it[k] = "" if it[k] is None else str(it[k]).strip()
        it["price"] = int(it["price"] or 0); it["rank"] = int(it["rank"] or 0)
        items.append(it)
    return items

def validate(items, slot_types, recipes):
    errs, seen = [], set()
    for it in items:
        w = f'{it["id"] or "(blank id)"}: '
        if not re.fullmatch(r"[a-z0-9][a-z0-9-]*", it["id"]): errs.append(w + "ItemID must be lowercase letters, digits, dashes")
        if it["id"] in seen: errs.append(w + "duplicate ItemID")
        seen.add(it["id"])
        if it["kind"] not in KINDS: errs.append(w + f'Kind "{it["kind"]}" is not one of {"/".join(KINDS)}')
        if it["kind"] == "Award" and it["price"] != 0: errs.append(w + "Awards are earned, never sold: Price must be 0")
        if it["price"] < 0: errs.append(w + "negative price")
        if it["slot"] in PAINT:
            v = it["sprite"]
            if it["slot"] == "paint-tiles":
                if v not in PATTERNS: errs.append(w + f"floor pattern must be one of {sorted(PATTERNS)}")
            elif v != "any" and not re.fullmatch(r"#[0-9A-Fa-f]{6}", v): errs.append(w + 'paint Sprite must be a #RRGGBB color, or "any" for the pick-any-color paint')
        elif it["slot"] not in slot_types: errs.append(w + f'unknown Slot "{it["slot"]}"')
        elif it["sprite"] not in recipes: errs.append(w + f'Sprite recipe "{it["sprite"]}" is not in scene.js RECIPES')
        if it["unlocks"] and it["kind"] != "Equipment": errs.append(w + "only Equipment can unlock case types")
    for pk in PAINT:
        if not any(i["slot"] == pk and i["price"] == 0 and i["kind"] != "Award" and i["sprite"] != "any" for i in items): errs.append(f"no free starter row for {pk}")
    if errs: sys.exit("Items tab problems:\n  " + "\n  ".join(errs))

def main(xlsx, out):
    items = read_items(xlsx)
    os.makedirs(os.path.join(out, "lab-sprites"), exist_ok=True)
    base = serve(os.path.dirname(os.path.abspath(__file__)))
    with sync_playwright() as p:
        b = p.chromium.launch(args=GL_ARGS); pg = b.new_page()
        errors = []; pg.on("pageerror", lambda e: errors.append(str(e)))
        pg.goto(base + "/bake.html"); pg.wait_for_function("window.READY", timeout=60000)
        meta = pg.evaluate("({v: api.sceneVersion, frame: api.frame, slots: api.slots, recipes: api.recipes})")
        slot_types = {s["type"] for s in meta["slots"]}
        validate(items, slot_types, set(meta["recipes"]))
        light = pg.evaluate("api.measure()")
        room = pg.evaluate("api.faces()")
        jobs = []
        for it in items:
            if it["slot"] in PAINT: continue
            for s in meta["slots"]:
                if s["type"] == it["slot"] and (it["sprite"], s["id"]) not in jobs: jobs.append((it["sprite"], s["id"]))
        sprites = {}
        for recipe, slot in jobs:
            url = pg.evaluate("([r,s,o]) => api.bakeItem(r,s,o)", [recipe, slot, light["shadowOpacity"]])
            im = Image.open(io.BytesIO(base64.b64decode(url.split(",", 1)[1]))).convert("RGBA")
            box = im.getchannel("A").point(lambda a: 255 if a > 2 else 0).getbbox()
            if not box: sys.exit(f"{recipe}@{slot} rendered nothing")
            crop = im.crop(box); fname = f"{recipe}__{slot}.png"
            crop.save(os.path.join(out, "lab-sprites", fname), optimize=True)
            sprites[f"{recipe}@{slot}"] = {"file": fname, "x": box[0], "y": box[1], "w": box[2] - box[0], "h": box[3] - box[1]}
            print(f"  baked {recipe:18s} @ {slot:12s} {box[2]-box[0]:4d}x{box[3]-box[1]:<4d} at ({box[0]},{box[1]})")
        thumbs = {}
        for recipe in dict.fromkeys(r for r, _ in jobs):
            slot = next(s for r, s in jobs if r == recipe)
            url = pg.evaluate("([r,s]) => api.bakeThumb(r,s)", [recipe, slot])
            im = Image.open(io.BytesIO(base64.b64decode(url.split(",", 1)[1]))).convert("RGBA")
            im = im.crop(im.getchannel("A").point(lambda a: 255 if a > 2 else 0).getbbox())
            im.thumbnail((THUMB, THUMB), Image.LANCZOS)
            im = im.quantize(colors=128, method=Image.FASTOCTREE)  # thumbs only; room sprites stay full RGBA
            fname = f"thumb__{recipe}.png"; im.save(os.path.join(out, "lab-sprites", fname), optimize=True)
            thumbs[recipe] = fname
        # framed award posters: canvas corners per poster spot (only if some item uses the frame)
        canvases = {}
        if any(it["sprite"] == "poster-frame" for it in items):
            for s in meta["slots"]:
                if s["type"] == "poster": canvases[s["id"]] = pg.evaluate("(s) => api.canvasFor(s)", s["id"])
        if errors: sys.exit("page errors: " + "; ".join(errors))
        b.close()

    # hit area per slot = union of every sprite that can go there (so an empty spot is still tappable)
    slots = []
    for s in meta["slots"]:
        boxes = [v for k, v in sprites.items() if k.endswith("@" + s["id"])]
        x0 = min(v["x"] for v in boxes); y0 = min(v["y"] for v in boxes)
        x1 = max(v["x"] + v["w"] for v in boxes); y1 = max(v["y"] + v["h"] for v in boxes)
        e = {"id": s["id"], "type": s["type"], "label": s["label"], "parent": s.get("parent"), "required": bool(s.get("required")),
             "hit": [x0, y0, x1 - x0, y1 - y0]}
        if s["id"] in canvases: e["canvas"] = canvases[s["id"]]; e["normal"] = "pz" if s.get("wall") == "back" else "px"
        slots.append(e)
    pts = [pt for f in room["faces"] for pt in f["points"]]
    xs = [p[0] for p in pts] + [v["x"] for v in sprites.values()] + [v["x"] + v["w"] for v in sprites.values()]
    ys = [p[1] for p in pts] + [v["y"] for v in sprites.values()] + [v["y"] + v["h"] for v in sprites.values()]
    pad = 16
    view = [int(min(xs)) - pad, int(min(ys)) - pad]; view += [int(max(xs)) + pad - view[0], int(max(ys)) + pad - view[1]]
    view = [max(0, view[0]), max(0, view[1]), min(meta["frame"]["w"], view[2]), min(meta["frame"]["h"], view[3])]
    manifest = {
        "manifestVersion": MANIFEST_VERSION, "sceneVersion": meta["v"], "generatedFrom": os.path.basename(xlsx),
        "note": "GENERATED by bake/bake.py from the Items tab. Do not hand-edit.",
        "frame": meta["frame"], "view": view,
        "room": {"faces": room["faces"], "tilesAfter": room["tilesAfter"], "tiles": room["tiles"],
                 "tileDarken": room["tileDarken"], "shade": light["shade"]},
        "shadowOpacity": light["shadowOpacity"],
        "slots": slots, "sprites": sprites, "thumbs": thumbs,
    }
    json.dump(manifest, open(os.path.join(out, "lab-manifest.json"), "w"), indent=1)
    json.dump(items, open(os.path.join(out, "lab-items.json"), "w"), indent=1, ensure_ascii=False)
    total = sum(os.path.getsize(os.path.join(out, "lab-sprites", f)) for f in os.listdir(os.path.join(out, "lab-sprites")))
    print(f"{len(items)} items, {len(sprites)} sprites, {total/1024:.0f} KB of PNG, view {view}")

if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
